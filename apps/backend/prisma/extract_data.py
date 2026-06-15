import json, sys

# Try to read the Excel file first (it's easier)
try:
    import openpyxl
    wb = openpyxl.load_workbook(r'c:\Users\z1309\Documents\xwechat_files\wxid_5pg1904cumf012_bc58\msg\file\2026-06\日常口语3000高频词.xlsx', read_only=True)
    ws = wb.active
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 100:  # First 100 rows
            rows.append(str(row))
        else:
            break
    with open(r'c:\Users\z1309\Desktop\work\speakGuild\apps\backend\prisma\extract_output.txt', 'w', encoding='utf-8') as f:
        f.write('=== Excel ===\n')
        f.write('\n'.join(rows))
except Exception as e:
    with open(r'c:\Users\z1309\Desktop\work\speakGuild\apps\backend\prisma\extract_output.txt', 'w', encoding='utf-8') as f:
        f.write(f'Excel error: {e}\n')

# Try to read the doc file
try:
    import docx
    doc = docx.Document(r'c:\Users\z1309\Documents\xwechat_files\wxid_5pg1904cumf012_bc58\msg\file\2026-06\常用英语500句.doc')
    paras = []
    for i, p in enumerate(doc.paragraphs):
        if i < 200:
            if p.text.strip():
                paras.append(p.text.strip())
        else:
            break
    with open(r'c:\Users\z1309\Desktop\work\speakGuild\apps\backend\prisma\extract_output.txt', 'a', encoding='utf-8') as f:
        f.write('\n\n=== DOC ===\n')
        f.write('\n'.join(paras))
except Exception as e:
    with open(r'c:\Users\z1309\Desktop\work\speakGuild\apps\backend\prisma\extract_output.txt', 'a', encoding='utf-8') as f:
        f.write(f'\n\nDoc error: {e}\n')

with open(r'c:\Users\z1309\Desktop\work\speakGuild\apps\backend\prisma\extract_output.txt', 'a', encoding='utf-8') as f:
    f.write('\n\nDone')
