import json

# Test what's available
try:
    import openpyxl
    print('openpyxl: OK')
    wb = openpyxl.load_workbook(r'c:\Users\z1309\Documents\xwechat_files\wxid_5pg1904cumf012_bc58\msg\file\2026-06\日常口语3000高频词.xlsx', read_only=True)
    ws = wb.active
    print(f'Sheet: {ws.title}')
    print(f'Dimensions: {ws.dimensions}')
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 20:
            print(f'Row {i}: {row}')
        else:
            break
except ImportError:
    print('openpyxl not available')
except Exception as e:
    print(f'openpyxl error: {e}')

try:
    import docx
    print('\ndocx: OK')
    doc = docx.Document(r'c:\Users\z1309\Documents\xwechat_files\wxid_5pg1904cumf012_bc58\msg\file\2026-06\常用英语500句.doc')
    count = 0
    for p in doc.paragraphs:
        if p.text.strip():
            print(f'Para: {p.text.strip()[:200]}')
            count += 1
            if count >= 30:
                break
except ImportError:
    print('\ndocx (python-docx): not available (old .doc format may need different library)')
except Exception as e:
    print(f'\ndocx error: {e}')
    print('The file might be old .doc format, not .docx')
